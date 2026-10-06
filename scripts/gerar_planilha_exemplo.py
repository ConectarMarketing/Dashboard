"""Gera data/exemplo.xlsx: planilha fictícia no mesmo layout do controle de mídia.

Uso: python3 scripts/gerar_planilha_exemplo.py   (requer: pip install openpyxl)
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HOJE = date(2026, 10, 6)
FIM = date(2026, 10, 31)

# cliente: (nome completo, verba mensal, [(plataforma, % da verba, [(métrica, verba, gasto, meta, resultado)])], anotações)
CLIENTES = {
    "Loja Exemplo": ("Loja Exemplo Ltda.", 3000, [
        ("Meta", 0.6, [
            ("Seguidores", 600, 140.5, 400, 120),
            ("Mensagens", 1200, 260.1, 90, 14),
        ]),
        ("Google", 0.4, [
            ("Search (cliques)", 1200, 230.9, 1500, 410),
        ]),
    ], ["Campanha de Black Friday começa dia 15"]),
    "Indústria B2B": ("Indústria Exemplo S.A.", 8000, [
        ("LinkedIn", 0.5, [
            ("Lead Gen - Decisores", 3000, 690.0, 40, 9),
            ("Conteúdo - Awareness", 1000, 205.3, 50000, 14200),
        ]),
        ("Google", 0.5, [
            ("Search (leads)", 4000, 980.4, 60, 11),
        ]),
    ], []),
    "Clínica": ("Clínica Exemplo", 1200, [
        ("Meta", 1.0, [
            ("Agendamentos (mensagens)", 900, 118.2, 60, 16),
            ("Alcance (regional)", 300, 31.0, 40000, 6100),
        ]),
    ], []),
}

AZUL = Font(color="1F4E9E")
ROXO = Font(color="7030A0")
NEGRITO = Font(bold=True)
CINZA = PatternFill("solid", fgColor="EDEDED")


def cabecalho(ws, linha, textos):
    for i, t in enumerate(textos, start=1):
        c = ws.cell(linha, i, t)
        c.font = NEGRITO
        c.fill = CINZA


wb = Workbook()
wb.remove(wb.active)
for aba, (nome, verba, plataformas, notas) in CLIENTES.items():
    ws = wb.create_sheet(aba)
    ws["A1"] = nome
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Geral"
    cabecalho(ws, 3, ["Dia atual", "Último dia do mês", "Dias restantes do mês",
                      "Valor de mídia mensal", "Valor de mídia gasto", "Valor de mídia restante"])
    ws["A4"], ws["B4"], ws["C4"], ws["D4"] = HOJE, FIM, "=B4-A4", verba
    ws["A4"].number_format = ws["B4"].number_format = "dd/mm/yyyy"
    ws["A4"].font = AZUL
    ws["D4"].font = ROXO
    linha = 5
    totais = []
    for plat, pct, metricas in plataformas:
        ws.cell(linha, 1, plat).font = NEGRITO
        cabecalho(ws, linha + 1, ["% de mídia mensal", "Valor de mídia mensal", "Valor de mídia diário",
                                  "Valor de mídia gasto mensal", "Valor de mídia restante",
                                  "Quanto gastar por dia restante"])
        v = linha + 2
        ini, fim = v + 2, v + 1 + len(metricas)
        ws.cell(v, 1, pct).font = ROXO
        ws.cell(v, 2, f"=D4*A{v}")
        ws.cell(v, 3, f"=B{v}/30.4")
        ws.cell(v, 4, f"=SUM(C{ini}:C{fim})")
        ws.cell(v, 5, f"=B{v}-D{v}")
        ws.cell(v, 6, f"=IFERROR(E{v}/C4,0)")
        totais.append(v)
        cabecalho(ws, v + 1, ["Métrica", "Valor de mídia", "Valor de mídia gasto", "Meta", "Resultado",
                              "Meta conquistada %"])
        for i, (m, verba_m, gasto, meta, res) in enumerate(metricas):
            r = ini + i
            ws.cell(r, 1, m)
            ws.cell(r, 2, verba_m).font = ROXO
            ws.cell(r, 3, gasto).font = AZUL
            ws.cell(r, 4, meta).font = ROXO
            ws.cell(r, 5, res).font = AZUL
            ws.cell(r, 6, f"=IFERROR(E{r}/D{r},0)").number_format = "0%"
        linha = fim + 2
    ws["E4"] = "=" + "+".join(f"D{v}" for v in totais)
    ws["F4"] = "=" + "+".join(f"E{v}" for v in totais)
    for t in notas:
        ws.cell(linha, 2, t)
        linha += 1
    linha += 1
    for t in ["Legenda:", "Azul: atualizar diariamente", "Roxo: atualizar mensalmente", "Preto: cálculo automático"]:
        ws.cell(linha, 1, t)
        linha += 1
    for col, larg in zip("ABCDEF", [34, 18, 18, 22, 20, 26]):
        ws.column_dimensions[col].width = larg

saida = Path(__file__).resolve().parent.parent / "data" / "exemplo.xlsx"
wb.save(saida)
print(f"Gerado: {saida}")
